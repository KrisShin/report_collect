"""Excel 附件生成。"""

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings

HEADERS = ["序号", "单位", "职务", "姓名", "电话", "邮箱", "提交时间"]


def build_submissions_xlsx(submissions: list, sheet_title: str = "报告信息") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for column, title in enumerate(HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    tz = ZoneInfo(settings.timezone)
    for index, item in enumerate(submissions, start=1):
        submitted_at = item.submitted_at
        if submitted_at and submitted_at.tzinfo is None:
            submitted_at = submitted_at.replace(tzinfo=ZoneInfo("UTC"))
        submitted_text = (
            submitted_at.astimezone(tz).strftime("%Y-%m-%d %H:%M")
            if submitted_at
            else ""
        )
        worksheet.append(
            [
                index,
                item.organization,
                item.position,
                item.name,
                item.phone,
                item.email,
                submitted_text,
            ]
        )
        worksheet.cell(row=index + 1, column=5).number_format = "@"  # 手机号按文本
        worksheet.cell(row=index + 1, column=6).number_format = "@"

    worksheet.freeze_panes = "A2"
    widths = [8, 28, 20, 14, 16, 30, 20]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def default_attachment_name(now: datetime | None = None) -> str:
    tz = ZoneInfo(settings.timezone)
    moment = now.astimezone(tz) if now else datetime.now(tz)
    return f"报告收集_{moment.strftime('%Y%m%d_%H%M%S')}.xlsx"

